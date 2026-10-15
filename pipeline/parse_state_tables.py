"""Read the state's other tables (Bergen rows only) into tidy CSV files.

- DirRatios.xlsx          -> data/processed/director_ratios_bergen.csv
- CoefficientDeviations.pdf -> data/processed/state_cod_bergen.csv
- GTRhistory.xlsx          -> data/processed/tax_rates_bergen.csv
- 2026AvgResStat.pdf       -> data/processed/avg_residential_bergen.csv
- 2025TEV.pdf, 2026TEV.pdf -> data/processed/tev_bergen.csv

Every table keys on district_code ("0201" ... "0270").
Run: .venv\\Scripts\\python.exe -m pipeline.parse_state_tables
"""

import re

import openpyxl
import pandas as pd
import pdfplumber

from pipeline import config


def code4(value):
    """Turn a code like 201 or "201" into the 4-digit text code "0201"."""
    return f"{int(value):04d}"


# ---------------------------------------------------------------- Excel tables

def read_wide_sheet(xlsx_path, year_row, first_data_row, county=config.COUNTY_NAME):
    """Read a state Excel sheet with one row per town and one column per year.

    Returns a long table: district_code, district_name, year, value.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    years = {i: int(v) for i, v in enumerate(rows[year_row]) if isinstance(v, (int, float))}
    out = []
    for row in rows[first_data_row:]:
        if not row or str(row[0]).strip().upper() != county.upper():
            continue
        for i, year in years.items():
            if row[i] is None:
                continue
            out.append({"district_code": code4(row[1]),
                        "district_name": str(row[2]).strip(),
                        "year": year, "value": float(row[i])})
    return pd.DataFrame(out)


def parse_director_ratios(path=config.DIR_RATIOS_XLSX):
    """Director's ratio history, long format, ratio in percent (2 decimals).

    The column year is the year the ratio was promulgated (October 1).
    Example: Allendale 2026 = 89.97, the same number as the 2027 Chapter 123
    average ratio.
    """
    df = read_wide_sheet(path, year_row=1, first_data_row=2)
    df["director_ratio"] = df.pop("value").round(2)
    df["source_file"] = "data/raw/state/DirRatios.xlsx"
    return df.sort_values(["district_code", "year"]).reset_index(drop=True)


def parse_tax_rates(path=config.GTR_XLSX):
    """General tax rates, long format, dollars per $100 of assessed value."""
    df = read_wide_sheet(path, year_row=0, first_data_row=4)
    df["general_tax_rate"] = df.pop("value").round(3)
    df["source_file"] = "data/raw/state/GTRhistory.xlsx"
    return df.sort_values(["district_code", "year"]).reset_index(drop=True)


# ---------------------------------------------------------------- Coefficients of deviation

COD_COLUMNS = ["general_coeff",
               "stratified_class1", "stratified_class2", "stratified_class4",
               "segmented_class1", "segmented_class2", "segmented_class4",
               "n_class1", "n_class2", "n_class4", "n_total"]
COD_CODE_LINE = re.compile(r"^(?P<name>.*?)\s*(?P<code>\d{4})\s+(?P<star>\*\s+)?"
                           r"(?P<year>20\d\d)\s+(?P<vals>.+)$")
COD_YEAR_LINE = re.compile(r"^(?P<star>\*\s+)?(?P<year>20\d\d)\s+(?P<vals>.+)$")
NAME_ONLY = re.compile(r"^[A-Za-z][A-Za-z .'-]*$")


def cod_values(text):
    """Split the 11 numbers after the year. "-" means blank."""
    parts = text.split()
    if len(parts) != len(COD_COLUMNS):
        raise ValueError(f"Expected {len(COD_COLUMNS)} values, got {len(parts)}: {text!r}")
    vals = [None if p == "-" else float(p.replace(",", "")) for p in parts]
    return dict(zip(COD_COLUMNS, vals))


def parse_cod_text(text, county=config.COUNTY_CODE, county_name=config.COUNTY_NAME):
    """Turn the CoefficientDeviations.pdf text into rows for one county."""
    rows, code, name, last_was_code = [], None, None, False
    for line in text.splitlines():
        line = line.strip()
        m = COD_CODE_LINE.match(line)
        if m and not COD_YEAR_LINE.match(line):
            code = m["code"]
            name = re.sub(rf"^{county_name}\s+", "", m["name"]).strip()
            year, star, vals, last_was_code = m["year"], m["star"], m["vals"], True
        elif (y := COD_YEAR_LINE.match(line)) and code:
            year, star, vals, last_was_code = y["year"], y["star"], y["vals"], False
        elif last_was_code and NAME_ONLY.match(line) and rows and rows[-1]["district_code"] == code:
            rows[-1]["district_name"] += " " + line  # name wrapped onto the next line
            name = rows[-1]["district_name"]
            continue
        else:
            last_was_code = False
            continue
        if not code.startswith(county):
            continue
        rows.append({"district_code": code, "district_name": name, "year": int(year),
                     "amended": star is not None, **cod_values(vals)})
    # A wrapped name is only known after its first row; copy it to the others.
    df = pd.DataFrame(rows)
    df["district_name"] = df.groupby("district_code")["district_name"].transform("last")
    return df


def parse_state_cod(path=config.COD_PDF):
    """State coefficients of deviation for Bergen, by town and year."""
    with pdfplumber.open(path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    df = parse_cod_text(text)
    for col in ["n_class1", "n_class2", "n_class4", "n_total"]:
        df[col] = df[col].astype("Int64")
    df["source_file"] = "data/raw/state/CoefficientDeviations.pdf"
    return df.sort_values(["district_code", "year"]).reset_index(drop=True)


# ---------------------------------------------------------------- Average residential stats

AVG_RES_HEADERS = {"Items": "line_items", "Assessment": "avg_assessment",
                   "Bill": "avg_tax_bill", "Sales": "n_sales", "Price": "avg_sale_price"}


def money(text):
    """Turn "$1,018,189.40" or "2,253" into a number."""
    return float(text.replace("$", "").replace(",", ""))


def parse_avg_res_page(page, county=config.COUNTY_CODE):
    """Read one page of 2026AvgResStat.pdf using each word's position.

    Numbers are matched to the header whose right edge is closest, so an
    empty column (Avg Tax Bill) stays empty instead of shifting values.
    """
    words = page.extract_words()
    header_top = next(w["top"] for w in words if w["text"] == "Items")
    header = {}
    for w in sorted(words, key=lambda w: w["x0"]):
        col = AVG_RES_HEADERS.get(w["text"])
        # "Sales" appears twice ("# of Sales", "Avg Sales Price"); keep the first.
        if abs(w["top"] - header_top) < 2 and col and col not in header:
            header[col] = w["x1"]
    lines = {}
    for w in words:
        lines.setdefault(round(w["top"]), []).append(w)
    rows = []
    for top in sorted(lines):
        ws = sorted(lines[top], key=lambda w: w["x0"])
        if not re.fullmatch(rf"{county}\d\d", ws[0]["text"]):
            continue
        row = {"district_code": ws[0]["text"]}
        name, nums = [], []
        for w in ws[1:]:
            (nums if re.fullmatch(r"\$?[\d,]+(\.\d+)?", w["text"]) else name).append(w)
        row["district_name"] = " ".join(w["text"] for w in name)
        for w in nums:
            col = min(header, key=lambda c: abs(header[c] - w["x1"]))
            row[col] = money(w["text"])
        rows.append(row)
    return rows


def parse_avg_residential(path=config.AVG_RES_PDF):
    """2026 average residential statistics for Bergen towns."""
    rows = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            rows += parse_avg_res_page(page)
    df = pd.DataFrame(rows)
    for col in AVG_RES_HEADERS.values():
        if col not in df:
            df[col] = None
    df = df[["district_code", "district_name", *AVG_RES_HEADERS.values()]]
    df["line_items"] = df["line_items"].astype("Int64")
    df["n_sales"] = df["n_sales"].astype("Int64")
    df["source_file"] = "data/raw/state/2026AvgResStat.pdf"
    return df.sort_values("district_code").reset_index(drop=True)


# ---------------------------------------------------------------- Tables of Equalized Valuations

# One town row: code, name (Tax Court changes marked "**"), then 6 numbers:
# aggregate assessed, average ratio, aggregate true value, class II railroad,
# personal property, equalized valuation.
TEV_ROW = re.compile(r"^(\d{4}) (.+?)(\*\*)? ([\d,]+) ([\d.]+) ([\d,]+) ([\d,]+) ([\d,]+) ([\d,]+)$")


def parse_tev_text(text, year, county=config.COUNTY_CODE):
    """Read the town rows for one county from the text of a Table of Equalized Valuations."""
    rows = []
    for line in text.splitlines():
        m = TEV_ROW.match(line.strip())
        if not m or not m[1].startswith(county):
            continue
        rows.append({
            "year": year, "district_code": m[1], "district_name": m[2].strip(),
            "agg_assessed": int(m[4].replace(",", "")), "average_ratio": float(m[5]),
            "agg_true_value": int(m[6].replace(",", "")), "class2_railroad": int(m[7].replace(",", "")),
            "personal_property": int(m[8].replace(",", "")), "equalized_valuation": int(m[9].replace(",", "")),
            "amended_by_tax_court": m[3] is not None,
        })
    return rows


def parse_tev(year, path=None):
    """Bergen rows of one year's Table of Equalized Valuations."""
    path = path or config.TEV_PDFS[year]
    with pdfplumber.open(path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    df = pd.DataFrame(parse_tev_text(text, year))
    df["source_file"] = f"data/raw/state/{path.name}"
    return df.sort_values("district_code").reset_index(drop=True)


# ---------------------------------------------------------------- run all

def find_2026_tax_rate_file():
    """Look for a separately saved 2026 tax rate file in data/raw/state/."""
    return sorted(p for p in config.RAW_STATE_DIR.iterdir()
                  if "2026" in p.name and "rate" in p.name.lower())


def main():
    """Parse every state table and print a short report."""
    dr = parse_director_ratios()
    dr.to_csv(config.DIR_RATIOS_CSV, index=False)
    print(f"Director's ratios: {dr.district_code.nunique()} towns, years {dr.year.min()}-{dr.year.max()}")

    cod = parse_state_cod()
    cod.to_csv(config.STATE_COD_CSV, index=False)
    print(f"State CoD: {cod.district_code.nunique()} towns, years {sorted(cod.year.unique())}, {len(cod)} rows")

    tr = parse_tax_rates()
    tr.to_csv(config.TAX_RATES_CSV, index=False)
    newest = tr.year.max()
    print(f"Tax rates: {tr.district_code.nunique()} towns, years {tr.year.min()}-{newest}; "
          f"newest year = {newest}")
    extra = find_2026_tax_rate_file()
    print(f"Separate 2026 tax rate file in data/raw/state/: {[p.name for p in extra] or 'none found'}")

    ar = parse_avg_residential()
    ar.to_csv(config.AVG_RES_CSV, index=False)
    print(f"Avg residential: {len(ar)} towns; avg_tax_bill filled for {ar.avg_tax_bill.notna().sum()}")

    tev = pd.concat([parse_tev(year) for year in sorted(config.TEV_PDFS)], ignore_index=True)
    tev.to_csv(config.TEV_CSV, index=False)
    print("Equalized valuations: " + ", ".join(f"{y} {n} towns" for y, n in tev.groupby("year").size().items()))


if __name__ == "__main__":
    main()
